from __future__ import annotations
import argparse
import hashlib
import html
import io
import json
import platform
from zipfile import BadZipFile
from urllib.parse import quote
from pathlib import Path
from datetime import datetime, timezone
import importlib.metadata
import numpy as np
import pandas as pd
import yaml
from .core import numeric, timestamps, time_flags, hampel, freeze_candidates, interpolate_short, compare


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda: f.read(1024*1024), b''): h.update(chunk)
    return h.hexdigest()


def json_text(value):
    """Strict JSON: unavailable statistics are null rather than NaN."""
    def scrub(v):
        if isinstance(v,dict):return {str(k):scrub(x) for k,x in v.items()}
        if isinstance(v,(list,tuple)):return [scrub(x) for x in v]
        if isinstance(v,(float,np.floating)):return float(v) if np.isfinite(v) else None
        if isinstance(v,np.integer):return int(v)
        return v
    return json.dumps(scrub(value),ensure_ascii=False,indent=2,allow_nan=False,default=str)


def validate_config(c):
    allowed = {'version','reader','time','columns','expected_dt','jitter_fraction','interpolation','mode_column','allowed_modes','max_rows'}
    if not isinstance(c, dict) or set(c)-allowed: raise ValueError('unknown configuration key')
    if not isinstance(c.get('version'), str) or not c['version']: raise ValueError('version is required')
    if not isinstance(c.get('columns'), dict) or not c['columns']: raise ValueError('columns is required')
    def keys(obj, allowed, name):
        if not isinstance(obj,dict) or set(obj)-set(allowed): raise ValueError('unknown key in '+name)
    keys(c.get('reader',{}), ['sheet','header','encoding','sep','decimal','thousands'], 'reader')
    keys(c.get('time',{}), ['column','format','timezone','excel_epoch'], 'time')
    if not c.get('time',{}).get('column'): raise ValueError('time.column required')
    from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
    try: ZoneInfo(c['time'].get('timezone','Asia/Tokyo'))
    except (ZoneInfoNotFoundError,TypeError) as exc: raise ValueError('unknown timezone') from exc
    if c['time'].get('excel_epoch') not in (None,'1900','1904'): raise ValueError('excel_epoch must be 1900 or 1904 as a string')
    for name, spec in c['columns'].items():
        keys(spec,['source','unit','target_unit','scale','offset','missing','lower','upper','window','k','floor','freeze_run','saturation'],name)
        if not isinstance(name,str) or not name or not spec.get('source'): raise ValueError('column source required')
        for key in ('scale','offset','lower','upper','k','floor'):
            if key in spec and (isinstance(spec[key],bool) or not isinstance(spec[key],(int,float)) or not np.isfinite(spec[key])): raise ValueError('finite number required: '+key)
        if spec.get('scale',1)==0: raise ValueError('scale cannot be zero')
        if spec.get('lower',-np.inf)>=spec.get('upper',np.inf): raise ValueError('lower must be smaller than upper')
        if spec.get('k',3.5)<=0 or spec.get('floor',0.01)<=0: raise ValueError('k and floor must be positive')
        w = spec.get('window',11)
        if type(w) is not int or w<3 or w%2!=1: raise ValueError('window must be an odd integer >=3')
        if type(spec.get('freeze_run',8)) is not int or spec.get('freeze_run',8)<2: raise ValueError('freeze_run >=2')
        if 'saturation' in spec and (not isinstance(spec['saturation'],list) or not all(isinstance(v,(int,float)) and np.isfinite(v) for v in spec['saturation'])): raise ValueError('saturation must be finite rails')
    if 'expected_dt' in c and (not isinstance(c['expected_dt'],(float,int)) or c['expected_dt']<=0 or not np.isfinite(c['expected_dt'])): raise ValueError('positive expected_dt required')
    if not 0 <= c.get('jitter_fraction',0.1) < 1: raise ValueError('jitter_fraction must be in [0,1)')
    if type(c.get('max_rows',200000)) is not int or not 1<=c.get('max_rows',200000)<=1000000: raise ValueError('max_rows must be 1..1000000')
    p = c.get('interpolation',{})
    keys(p,['enabled','max_gap_seconds','reason','approved_by'], 'interpolation')
    if type(p.get('enabled',False)) is not bool: raise ValueError('enabled must be Boolean')
    if p.get('enabled',False):
        if not p.get('reason') or not p.get('approved_by') or not isinstance(p.get('max_gap_seconds'),(float,int)) or not np.isfinite(p['max_gap_seconds']) or p['max_gap_seconds']<=0:
            raise ValueError('interpolation requires reason, approved_by and positive max_gap_seconds')
    if 'allowed_modes' in c and (not c.get('mode_column') or not isinstance(c['allowed_modes'],list)): raise ValueError('allowed_modes requires mode_column')
    return c


def read_input(path,c):
    r = c.get('reader',{}); header = r.get('header',0)
    if type(header) is not int or header<0: raise ValueError('reader.header must be a nonnegative integer')
    limit = c.get('max_rows',200000)
    def unique_headers(values):
        nonempty=[str(x) for x in values if x is not None and str(x).strip()!='']
        if len(nonempty)!=len(set(nonempty)):raise ValueError('duplicate column headers; choose an explicit input adapter')
    if path.suffix.lower()=='.csv':
        import csv
        sep=r.get('sep',',')
        if not isinstance(sep,str) or len(sep)!=1:raise ValueError('CSV separator must be one character')
        with path.open(encoding=r.get('encoding','utf-8-sig'),newline='') as f:
            rows=csv.reader(f,delimiter=sep)
            for _ in range(header+1):
                names=next(rows,None)
                if names is None:raise ValueError('missing CSV header')
            unique_headers(names)
        # Strings and empty strings survive ingestion; parsing happens separately.
        frames={'CSV':pd.read_csv(path,header=header,dtype=str,keep_default_na=False,skip_blank_lines=False,encoding=r.get('encoding','utf-8-sig'),sep=sep,nrows=limit+1)}
    elif path.suffix.lower()=='.xlsx':
        frames = pd.read_excel(path,sheet_name=r.get('sheet',0),header=header,dtype=object,keep_default_na=False,nrows=limit+1,engine='openpyxl')
        if not isinstance(frames,dict): frames={str(r.get('sheet',0)):frames}
        from openpyxl import load_workbook
        wb=load_workbook(path,read_only=True,data_only=True)
        try:
            selected=r.get('sheet',0)
            for sheet in frames:
                ws=wb.worksheets[selected] if type(selected) is int else wb.worksheets[sheet] if type(sheet) is int else wb[sheet]
                unique_headers(next(ws.iter_rows(min_row=header+1,max_row=header+1,values_only=True),()))
        finally:wb.close()
    else: raise ValueError('only .csv and .xlsx are supported')
    for sheet,df in frames.items():
        if len(df)>limit: raise ValueError('max_rows exceeded; split input using documented streaming workflow')
        if df.empty: raise ValueError('empty input')
    return frames


def process_frame(df,c,source='',sheet=''):
    c=validate_config(c); df=df.reset_index(drop=True)
    if df.empty: raise ValueError('empty input')
    required=[c['time']['column']]+[s['source'] for s in c['columns'].values()]
    if c.get('mode_column'): required.append(c['mode_column'])
    if any(x not in df.columns for x in required): raise ValueError('missing required columns: '+str(set(required)-set(df.columns)))
    header=c.get('reader',{}).get('header',0)
    provenance=pd.DataFrame({'source_file':source,'sheet':sheet,'source_row':np.arange(len(df))+header+2,'source_row_kind':'logical_csv_record' if sheet=='CSV' else 'worksheet_row' if sheet else 'input_table_row','record_index':np.arange(len(df))})
    time_spec={k:v for k,v in c['time'].items() if k!='column'}
    t=timestamps(df[c['time']['column']],**time_spec)
    tf=time_flags(t,c.get('expected_dt'),c.get('jitter_fraction',0.1))
    out=provenance.copy(); out['time_raw']=df[c['time']['column']]; out['time_utc']=t.map(lambda x: x.isoformat() if pd.notna(x) else '')
    flags=pd.concat([provenance,tf],axis=1)
    warnings=[]; logs=[]; summaries={}
    if tf[['time_invalid','time_duplicate','time_reverse']].any().any(): warnings.append('TIME_AXIS_REVIEW: interpolation disabled; original acquisition order preserved')
    modes=df[c['mode_column']].astype(str) if c.get('mode_column') else pd.Series('all',index=df.index)
    # Segment identity changes at every mode transition, including A -> B -> A.
    breaks=modes.ne(modes.shift()) | tf['time_invalid'] | tf['time_reverse'] | tf['time_duplicate'] | tf['interval_irregular']
    segments=breaks.cumsum()
    r=c.get('reader',{})
    for name,s in c['columns'].items():
        raw=df[s['source']]
        v,missing,invalid=numeric(raw,unit=s.get('unit',''),decimal=r.get('decimal','.'),thousands=r.get('thousands',''),missing=s.get('missing',['','NA','---','欠測']))
        with np.errstate(over='ignore',invalid='ignore'):
            v=v*s.get('scale',1)+s.get('offset',0)
        overflow=v.notna() & ~np.isfinite(v)
        invalid=invalid|overflow;v=v.where(np.isfinite(v))
        h=pd.DataFrame({'candidate':False,'assessed':False},index=v.index)
        freeze=pd.Series(False,index=v.index)
        for _,idx in segments.groupby(segments).groups.items():
            hr=hampel(v.loc[idx].to_numpy(),window=s.get('window',11),k=s.get('k',3.5),floor=s.get('floor',0.01))
            h.loc[idx,['candidate','assessed']]=hr[['candidate','assessed']].to_numpy()
            freeze.loc[idx]=freeze_candidates(v.loc[idx].to_numpy(),s.get('freeze_run',8)).to_numpy()
        bounds=v.lt(s.get('lower',-np.inf)) | v.gt(s.get('upper',np.inf))
        saturation=v.isin(s.get('saturation',[]))
        # Parenthesize the numeric sign test separately (Boolean precedence).
        spike=h['candidate'] & v.shift().notna() & v.shift(-1).notna() & (((v-v.shift())*(v.shift(-1)-v))<0)
        if v.notna().sum()==0: warnings.append(name+': NO_VALID_VALUES')
        elif not h['assessed'].any(): warnings.append(name+': HAMPEL_NOT_ASSESSED')
        y=v.copy(); filled=pd.Series(False,index=v.index)
        interp=c.get('interpolation',{})
        if interp.get('enabled',False) and not tf[['time_invalid','time_duplicate','time_reverse']].any().any():
            eligible=missing & modes.isin(c.get('allowed_modes',modes.unique().tolist()))
            anchors=~(h['candidate']|bounds|saturation|freeze)
            y,filled=interpolate_short(t,v,max_gap_seconds=interp['max_gap_seconds'],eligible=eligible,groups=segments,anchors_ok=anchors)
            for i in np.flatnonzero(filled):
                logs.append({'source_file':source,'sheet':sheet,'source_row':int(provenance.loc[i,'source_row']),'column':name,'rule':'short_gap_linear','before':str(raw.iloc[i]),'after':float(y.iloc[i]),'reason':interp['reason'],'approved_by':interp['approved_by']})
        out[name+'__raw']=raw; out[name+'__numeric']=v; out[name+'__analysis']=y
        out[name+'__origin']=np.select([filled,invalid,missing],['interpolated','invalid_token','missing'],default='observed')
        out[name+'__unit']=s.get('target_unit',s.get('unit',''))
        for label,mask in {'missing':missing,'invalid_token':invalid,'bounds_candidate':bounds,'outlier_candidate':h['candidate'],'hampel_assessed':h['assessed'],'spike_candidate':spike,'freeze_candidate':freeze,'saturation_candidate':saturation,'interpolated':filled}.items():
            flags[name+'__'+label]=mask
        summaries[name]=compare(v,y)
    candidate_cols=[x for x in flags if x.endswith('_candidate')]
    problem_cols=candidate_cols+[x for x in flags if x.endswith('__missing') or x.endswith('__invalid_token')]+['time_invalid','time_duplicate','time_reverse','interval_irregular']
    status='REVIEW_REQUIRED' if flags[problem_cols].any().any() or warnings else 'NO_RULE_VIOLATION_DETECTED'
    metrics={}
    for name in c['columns']:
        assessed=int(flags[name+'__hampel_assessed'].sum())
        candidates=int(flags[name+'__outlier_candidate'].sum())
        metrics[name]={'denominator_received_rows':len(df),'missing_fraction_received_rows':float(flags[name+'__missing'].mean()),'invalid_fraction_received_rows':float(flags[name+'__invalid_token'].mean()),'hampel_assessed_n':assessed,'hampel_unassessed_n':len(df)-assessed,'candidate_fraction_assessed':candidates/assessed if assessed else None}
    return out,flags,logs,{'status':status,'rows':len(df),'warnings':warnings,'statistics':summaries,'quality_metrics':metrics,'flag_counts':{x:int(flags[x].sum()) for x in flags if flags[x].dtype==bool}}


def excel_safe(df,path,sheet_name='data'):
    with pd.ExcelWriter(path,engine='openpyxl') as w:
        df.to_excel(w,index=False,sheet_name=sheet_name)
        ws=w.sheets[sheet_name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        from openpyxl.styles import Font, PatternFill
        for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24576B')
        from openpyxl.utils import get_column_letter
        for i in range(1,len(df.columns)+1):ws.column_dimensions[get_column_letter(i)].width=24
        # Prevent a raw string beginning '=' from becoming a workbook formula.
        for row in ws:
            for cell in row:
                if cell.data_type=='f':cell.data_type='s'


def svg_comparison(out,columns):
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt,font_manager
    import importlib.util
    spec=importlib.util.find_spec('japanize_matplotlib')
    if spec:
        font_manager.fontManager.addfont(str(Path(spec.origin).parent/'fonts/ipaexg.ttf'))
        plt.rcParams['font.family']='IPAexGothic'
    plt.rcParams.update({'svg.fonttype':'path','axes.unicode_minus':False})
    fig,axs=plt.subplots(len(columns),1,figsize=(10,3*len(columns)),squeeze=False)
    for ax,name in zip(axs[:,0],columns):
        # Draw all values: decimating a quality plot can hide a narrow peak.
        ax.plot(out[name+'__numeric'],color='#596a7a',lw=1,label='元の数値')
        ax.plot(out[name+'__analysis'],'--',color='#b04924',lw=1,label='解析用')
        mask=out[name+'__origin'].eq('interpolated')
        ax.scatter(out.index[mask],out.loc[mask,name+'__analysis'],marker='x',color='#b04924',label='補間値')
        ax.set(xlabel='記録番号（取得順）',ylabel=f"{name} [{out[name+'__unit'].iloc[0]}]");ax.legend();ax.grid(alpha=.2)
    fig.tight_layout();f=io.StringIO();fig.savefig(f,format='svg');plt.close(fig)
    return f.getvalue()[f.getvalue().index('<svg'):]


def write_result(folder,out,flags,logs,summary,c,meta):
    folder=Path(folder); folder.mkdir(parents=True,exist_ok=False)
    excel_safe(out,folder/'clean_data.xlsx');excel_safe(flags,folder/'quality_flags.xlsx','flags')
    logcols=['run_utc','source_sha256','config_sha256','config_version','source_file','sheet','source_row','column','rule','before','after','reason','approved_by']
    for row in logs: row.update({k:meta[k] for k in logcols[:4]})
    if not logs:logs=[{**{k:meta[k] for k in logcols[:4]},'source_file':out.source_file.iloc[0],'sheet':out.sheet.iloc[0],'rule':'inspection_only','reason':'No measurement values changed'}]
    pd.DataFrame(logs).reindex(columns=logcols).to_csv(folder/'processing_log.csv',index=False,encoding='utf-8-sig')
    (folder/'run_manifest.json').write_text(json_text({'summary':summary,'environment':meta,'effective_rules':c}),encoding='utf-8')
    boolcols=[x for x in flags if flags[x].dtype==bool and not x.endswith('__hampel_assessed')]
    candidates=flags.loc[flags[boolcols].any(axis=1)].head(200)
    body=f"""<!doctype html><html lang="ja"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>測定データ品質レポート</title><style>body{{max-width:1100px;margin:auto;padding:2rem;font:16px/1.8 sans-serif}}svg{{max-width:100%;height:auto}}pre{{white-space:pre-wrap;overflow-wrap:anywhere}}.table{{overflow:auto}}table{{border-collapse:collapse;font-size:12px}}td,th{{border:1px solid #ccc;padding:5px}}@media print{{.table{{overflow:visible}}}}</style><h1>測定データ品質レポート</h1><p>判定：{html.escape(summary['status'])}。フラグは原因を確定しません。未検査項目を含み、合格保証ではありません。</p><h2>品質指標・警告</h2><pre>{html.escape(json.dumps(summary,ensure_ascii=False,indent=2,default=str))}</pre><h2>処理前後</h2>{svg_comparison(out,c['columns'])}<h2>要確認行（先頭200行。全件はquality_flags.xlsx）</h2><div class="table">{candidates.to_html(index=False,escape=True)}</div><h2>適用ルール</h2><pre>{html.escape(yaml.safe_dump(c,allow_unicode=True,sort_keys=False))}</pre><h2>追跡情報</h2><pre>{html.escape(json.dumps(meta,ensure_ascii=False,indent=2))}</pre></html>"""
    (folder/'quality_report.html').write_text(body,encoding='utf-8')


def main(argv=None):
    p=argparse.ArgumentParser(description='Preserve data and detect quality candidates; all thresholds are project-specific.')
    p.add_argument('--input',nargs='+',required=True);p.add_argument('--config',required=True);p.add_argument('--output',required=True)
    args=p.parse_args(argv)
    cfgpath=Path(args.config); c=validate_config(yaml.safe_load(cfgpath.read_text(encoding='utf-8')))
    paths=[]
    for item in args.input:
        path=Path(item)
        paths.extend(sorted(x for x in path.iterdir() if x.suffix.lower() in ('.csv','.xlsx')) if path.is_dir() else [path])
    paths=list(dict.fromkeys(x.resolve() for x in paths))
    if not paths:raise ValueError('no input files')
    root=Path(args.output).resolve()
    if root.exists() and any(root.iterdir()):raise ValueError('output must be a new or empty directory; previous results are never overwritten')
    if any(root==x or root in x.parents for x in paths+[cfgpath.resolve()]):raise ValueError('input/config cannot be inside output')
    root.mkdir(parents=True,exist_ok=True)
    errors=[];results=[];run=datetime.now(timezone.utc).isoformat()
    for index,path in enumerate(paths):
        try:
            frames=read_input(path,c)
            for sheet_index,(sheet,frame) in enumerate(frames.items()):
                out,flags,logs,summary=process_frame(frame,c,str(path),sheet)
                meta={'run_utc':run,'source_sha256':sha256(path),'config_sha256':sha256(cfgpath),'config_version':c['version'],'python':platform.python_version(),'packages':{name:importlib.metadata.version(name) for name in ['numpy','pandas','scipy','matplotlib','openpyxl','PyYAML']}}
                dest=root/f'{index:03d}_{sheet_index:02d}_{path.stem}'
                write_result(dest,out,flags,logs,summary,c,meta)
                results.append({'file':str(path),'sheet':sheet,'folder':dest.name,'status':summary['status']})
        except (ValueError,OSError,KeyError,TypeError,BadZipFile) as exc:
            errors.append({'file':str(path),'error':str(exc)})
    (root/'batch_summary.json').write_text(json.dumps({'results':results,'errors':errors},ensure_ascii=False,indent=2),encoding='utf-8')
    links=''.join(f'<li><a href="{quote(x["folder"],safe="")}/quality_report.html">{html.escape(x["file"])} / {html.escape(x["sheet"])}</a> {x["status"]}</li>' for x in results)
    (root/'index.html').write_text('<!doctype html><html lang="ja"><meta charset="utf-8"><title>処理一覧</title><h1>処理一覧</h1><ul>'+links+'</ul><h2>エラー（原本は移動しません）</h2><pre>'+html.escape(json.dumps(errors,ensure_ascii=False,indent=2))+'</pre></html>',encoding='utf-8')
    print(json.dumps({'processed':len(results),'errors':len(errors)},ensure_ascii=False))
    return 2 if errors else 0
