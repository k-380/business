"""Recreate ten synthetic engineering datasets; no real test data is used."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.utils.datetime import to_excel

ROOT=Path(__file__).resolve().parents[2]
SEED=20261028


def main():
    rng=np.random.default_rng(SEED)
    raw=ROOT/'datasets/raw'; truth=ROOT/'datasets/expected'
    raw.mkdir(parents=True,exist_ok=True);truth.mkdir(parents=True,exist_ok=True)
    manifests=[]
    def save(name,t,y,observed,unit,events,model,extra=None,dt=None,truth_time=None):
        time=pd.Timestamp('2026-01-01T09:00:00+09:00')+pd.to_timedelta(t,unit='s')
        frame=pd.DataFrame({'時刻':time.astype(str),'測定値':observed})
        if extra:
            for k,v in extra.items():frame[k]=v
        frame.to_csv(raw/(name+'.csv'),index=False,encoding='utf-8-sig',float_format='%.10g')
        pd.DataFrame({'sample_id':np.arange(len(t)),'physical_time_s':t if truth_time is None else truth_time,'latent_signal':y}).to_csv(truth/(name+'_truth.csv'),index=False,float_format='%.12g')
        c={'version':'teaching-1','time':{'column':'時刻','format':'ISO8601','timezone':'Asia/Tokyo'},'columns':{'value':{'source':'測定値','unit':unit,'window':11,'k':3.5,'floor':0.01,'freeze_run':8}},'interpolation':{'enabled':False},'max_rows':200000}
        if dt:c['expected_dt']=dt
        if name=='04_current':c['columns']['value']['saturation']=[-12,12]
        if name=='02_pressure':c['columns']['value'].update(lower=0,upper=2.5)
        if extra and 'mode' in extra:c.update(mode_column='mode',allowed_modes=['run'])
        (ROOT/'configs'/f'{name}.yaml').write_text(yaml.safe_dump(c,allow_unicode=True,sort_keys=False),encoding='utf-8')
        manifests.append({'id':name,'seed':SEED,'n':len(t),'unit':unit,'model':model,'events':events,'index_convention':'zero-based, inclusive start, exclusive stop','raw_file':name+'.csv','truth_file':name+'_truth.csv','note':'Synthetic; latent_signal excludes measurement noise. Truth is unavailable in real tests.'})
        return frame,c
    t=np.arange(1440)*60.
    y=25+35*(1-np.exp(-t/1200))+0.8*np.sin(2*np.pi*t/86400)
    obs=y+rng.normal(0,.06,len(t))+np.maximum(t-43200,0)*.00003
    obs[120:123]=np.nan;obs[600:630]=np.nan
    save('01_temperature',t,y,obs,'degC',[{'start':120,'stop':123,'kind':'missing_short'},{'start':600,'stop':630,'kind':'missing_long'},{'start':721,'stop':1440,'kind':'sensor_drift','slope_per_second':.00003}], '25+35(1-exp(-t/1200))+0.8sin(2πt/86400); τ=1200 s; noise σ=.06 degC',dt=60)
    t=np.arange(600)*.1;y=1.4+.04*np.sin(2*np.pi*.3*t)+.55*np.exp(-((t-35)/.45)**2)
    obs=y+rng.normal(0,.004,len(t));obs[[90,270]]=[3.8,-.8]
    save('02_pressure',t,y,obs,'MPa',[{'start':90,'stop':91,'kind':'measurement_spike'},{'start':270,'stop':271,'kind':'measurement_spike'},{'start':338,'stop':363,'kind':'real_pressure_pulse','action':'preserve'}],'1.4+.04sin(2π .3t)+.55exp(-((t-35)/.45)^2); true pressure pulse',dt=.1)
    t=np.arange(4000)/1000;y=2*np.sin(2*np.pi*40*t)+.6*np.sin(2*np.pi*130*t)
    obs=y+.3*np.sin(2*np.pi*370*t)+rng.normal(0,.03,len(t))
    save('03_vibration',t,y,obs,'m/s2',[{'start':0,'stop':4000,'kind':'high_frequency_interference','frequency_hz':370}], '40 Hz structural motion +130 Hz component; 370 Hz interference; fs=1000 Hz',dt=.001)
    t=np.arange(4000)/2000;envelope=5+13*np.exp(-t/.08);y=envelope*np.sin(2*np.pi*50*t)
    obs=np.clip(y+rng.normal(0,.02,len(t)),-12,12)
    clipped=np.flatnonzero(np.abs(y)>12)
    save('04_current',t,y,obs,'A',[{'indices':clipped.tolist(),'kind':'saturation','rails':[-12,12]}], '(5+13exp(-t/.08))sin(2π50t); ADC limit ±12 A',dt=.0005)
    t=np.arange(500)*.1;y=np.exp(-((t-10)/.3)**2)+.6*np.exp(-((t-30)/.5)**2)
    obs=y+rng.normal(0,.003,len(t));clock=t*1.0002+.7
    save('05_sync',clock,y,obs,'V',[{'start':0,'stop':500,'kind':'clock_error','offset_s':.7,'scale':1.0002}], 'Two common injected trigger pulses at 10 s and 30 s; logger B clock tB=1.0002t+.7',extra={'reference_time_s':t,'reference_trigger':y},dt=.1,truth_time=t)
    t=np.arange(1000)*3600.;y=2+.0006*np.arange(1000)
    obs=y+.0008*np.maximum(np.arange(1000)-500,0)+rng.normal(0,.01,len(t))
    save('06_durability',t,y,obs,'mm/s',[{'start':0,'stop':1000,'kind':'real_degradation','slope_per_hour':.0006},{'start':501,'stop':1000,'kind':'sensor_drift','slope_per_hour':.0008}], 'RMS vibration grows .0006 mm/s per hour; sensor drift after 500 h; independent reference available',extra={'reference':y+rng.normal(0,.005,len(t))},dt=3600)
    t=np.arange(600)*10.;mode=np.where((t>=2000)&(t<3000),'stop','run');y=30+20*(1-np.exp(-t/400))
    y[200:300]=30+20*np.exp(-(t[200:300]-2000)/300)
    y[300:]=30+20*(1-np.exp(-(t[300:]-3000)/400))
    obs=y+rng.normal(0,.05,len(t));obs[205:295]=np.nan;obs[400:402]=np.nan
    save('07_modes',t,y,obs,'degC',[{'start':200,'stop':300,'kind':'stopped_operation'},{'start':205,'stop':295,'kind':'planned_no_measurement'},{'start':400,'stop':402,'kind':'missing_short'}], 'First-order heating/cooling with stop at 2000 s and restart at 3000 s',extra={'mode':mode},dt=10)
    t=np.arange(200,dtype=float);y=20+.02*t;clock=t.copy();clock[100:]-=80;clock[50]=clock[49]
    save('08_restart',clock,y,y,'degC',[{'start':50,'stop':51,'kind':'duplicate_timestamp'},{'start':100,'stop':101,'kind':'clock_reset','offset_s':-80}], 'Physical temperature slope .02 degC/s; acquisition sequence monotonic, recorded clock reset',extra={'sequence':np.arange(200)},dt=1,truth_time=t)
    t=np.arange(400,dtype=float);y=40+5*(1-np.exp(-t/50));obs=y+rng.normal(0,.03,len(t));obs[200:220]=obs[199]
    save('09_prototypes',t,y,obs,'degC',[{'start':199,'stop':220,'kind':'sensor_freeze'}], 'Prototype A temperature approaches 45 degC; prototype B +.8 degC, alternate sampling',extra={'prototype_B':np.where(np.arange(400)%2==0,y+.8,np.nan)},dt=1)
    t=np.arange(40)*30.;y=25+.01*t;obs=y.astype(object);obs[4]='---';obs[8]='故障';obs[10]=f'{y[10]} degC';obs[15]=''
    frame,c=save('10_irregular_excel',t,y,obs,'degC',[{'start':4,'stop':5,'kind':'missing_code'},{'start':8,'stop':9,'kind':'invalid_token'},{'start':10,'stop':11,'kind':'unit_suffix'},{'start':15,'stop':16,'kind':'blank_cell'}], 'Linear ramp .01 degC/s; irregular workbook has title merge, header metadata, blank row, serial dates',dt=30)
    wb=Workbook();ws=wb.active;ws.title='測定'
    ws.merge_cells('A1:C1');ws['A1']='模擬データ：温度試験（実測ではない）';ws.append(['ロガー','SYNTHETIC-01']);ws.append(['単位','degC','運転条件'])
    ws.append(['時刻','測定値','条件'])
    for i,row in frame.iterrows():
        ws.append([to_excel(pd.Timestamp(row['時刻']).tz_localize(None).to_pydatetime()),None if row['測定値']=='' else row['測定値'],'加熱' if i<20 else '保持'])
    ws.insert_rows(18)  # Deliberate blank record inside the table; never silently drop.
    ws.column_dimensions['A'].width=25;ws.column_dimensions['B'].width=20;ws.column_dimensions['C'].width=20
    ws.freeze_panes='A5';meta=wb.create_sheet('条件履歴');meta.append(['Excel行','変更']);meta.append([26,'保持へ切替（模擬）'])
    wb.save(raw/'10_irregular_excel.xlsx')
    c['reader']={'sheet':'測定','header':3};c['time']['excel_epoch']='1900'
    (ROOT/'configs/10_workbook.yaml').write_text(yaml.safe_dump(c,allow_unicode=True,sort_keys=False),encoding='utf-8')
    # Locale/encoding and split-file variants are deliberately separate inputs.
    pd.DataFrame({'時刻':['2026-01-01 09:00:00'],'測定値':['1.234,5']}).to_csv(raw/'locale_cp932.csv',index=False,encoding='cp932',sep=';')
    a=pd.read_csv(raw/'09_prototypes.csv');a.iloc[:210].to_csv(raw/'prototype_part1.csv',index=False);a.iloc[200:].assign(追加電流_A=1.2).to_csv(raw/'prototype_part2.csv',index=False)
    (ROOT/'datasets/manifest.json').write_text(json.dumps({'seed':SEED,'datasets':manifests,'workbook_note':'CSV sample i maps to Excel row i+5 before i=13, i+6 afterwards. Excel row 18 is blank; source_row preserves it.'},ensure_ascii=False,indent=2),encoding='utf-8')
    print('Generated 10 synthetic datasets, workbook, locale CSV and split-file examples.')

if __name__=='__main__':main()
