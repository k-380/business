from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
sys.path.insert(0,str(ROOT/'examples'))
import numpy as np
import pandas as pd
from common import load,truth,seconds,plot
from tsclean.core import numeric,timestamps,time_flags,hampel,freeze_candidates,interpolate_short,compare

from openpyxl import load_workbook
p=ROOT/'datasets/raw/10_irregular_excel.xlsx'
wb=load_workbook(p,data_only=False)
print('シート:',wb.sheetnames,'結合:',list(wb['測定'].merged_cells.ranges),'基準日:',wb.epoch)
preview=pd.read_excel(p,sheet_name='測定',header=None,nrows=6)
print(preview.to_string(index=False,header=False))
df=pd.read_excel(p,sheet_name='測定',header=3,keep_default_na=False)
x,missing,bad=numeric(df['測定値'],unit='degC')
print('読み込み行数:',len(df),'完全空白行:',int(df.eq('').all(axis=1).sum()),'変換不能:',int(bad.sum()))
plot(2,np.arange(len(df))+5,{'数値化可能な値':x},'温度 [°C]',xlabel='Excelの行番号')
